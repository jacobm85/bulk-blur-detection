"""Measure how well the sharpness score separates blurry from sharp photos.

Uses the CERTH Image Blur Dataset (https://mklab.iti.gr/results/certh-image-blur-dataset/):
download it, extract EvaluationSet/NaturalBlurSet and NaturalBlurSet.xlsx, then run

    pip install openpyxl
    python tools/evaluate_certh.py path/to/CERTH_ImageBlurDataset/EvaluationSet

It prints the AUC and, for a range of thresholds, how many blurry photos are found and how many
sharp photos would be flagged by mistake. Use it to check every change to scoring.py.
"""
import argparse
import os
import sys
from concurrent.futures import ProcessPoolExecutor

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scoring import DEFAULT_EYE_THRESHOLD, DEFAULT_THRESHOLD, score_file  # noqa: E402


def load_labels(xlsx):
    rows = load_workbook(xlsx, read_only=True).active.iter_rows(min_row=2, values_only=True)
    return {str(name).strip().lower(): label == 1 for name, label, *_ in rows if name}


def auc(scores, blurry):
    """Probability that a random blurry photo scores lower than a random sharp one."""
    ranks = np.argsort(np.argsort(scores)) + 1
    n1, n0 = blurry.sum(), (~blurry).sum()
    return 1 - (ranks[blurry].sum() - n1 * (n1 + 1) / 2) / (n1 * n0)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('evaluation_set', help='Path to CERTH_ImageBlurDataset/EvaluationSet')
    args = parser.parse_args()

    labels = load_labels(os.path.join(args.evaluation_set, 'NaturalBlurSet.xlsx'))
    folder = os.path.join(args.evaluation_set, 'NaturalBlurSet')
    files = sorted(f for f in os.listdir(folder) if os.path.splitext(f)[0].lower() in labels)
    with ProcessPoolExecutor() as pool:
        results = list(pool.map(score_file, [os.path.join(folder, f) for f in files], chunksize=8))

    ok = [(f, r) for f, r in zip(files, results) if r.get('error') is None]
    scores = np.array([r['score'] for _, r in ok])
    blurry = np.array([labels[os.path.splitext(f)[0].lower()] for f, _ in ok])
    print(f'{len(ok)} images ({blurry.sum()} blurry, {(~blurry).sum()} sharp), '
          f'{len(files) - len(ok)} failed to load')
    print(f'AUC {auc(scores, blurry):.3f}  (1.0 = perfect, 0.5 = coin flip)\n')
    print(f'{"threshold":>10}  {"blurry found":>14}  {"sharp flagged":>14}  {"precision":>9}')
    for t in sorted({50, 100, 150, 200, 250, 300, 400, 600, DEFAULT_THRESHOLD}):
        flagged = scores < t
        tp, fp = (flagged & blurry).sum(), (flagged & ~blurry).sum()
        marker = '  <- default' if t == DEFAULT_THRESHOLD else ''
        print(f'{t:>10}  {tp / blurry.sum():>14.1%}  {fp / (~blurry).sum():>14.1%}  '
              f'{tp / max(tp + fp, 1):>9.1%}{marker}')

    # Photos with a clear face are judged on the eyes instead (few in CERTH, so only a rough check)
    eyes = np.array([r['eye_score'] if r['eye_score'] is not None else np.nan for _, r in ok])
    has_eyes = ~np.isnan(eyes)
    flagged = np.where(has_eyes, eyes < DEFAULT_EYE_THRESHOLD, scores < DEFAULT_THRESHOLD)
    tp, fp = (flagged & blurry).sum(), (flagged & ~blurry).sum()
    print(f'\n{has_eyes.sum()} photos judged on the eyes ({(has_eyes & blurry).sum()} blurry). With the default '
          f'thresholds ({DEFAULT_THRESHOLD}, eyes {DEFAULT_EYE_THRESHOLD}): {tp / blurry.sum():.1%} of the blurry '
          f'photos found, {fp / (~blurry).sum():.1%} of the sharp ones flagged')
    for (name, r), is_blurry in sorted(((x, b) for x, b in zip(ok, blurry) if x[1]['eye_score'] is not None),
                                       key=lambda x: x[0][1]['eye_score']):
        print(f'    {name:<20} {"blurry" if is_blurry else "sharp":<7} eyes {r["eye_score"]:7.1f}  score {r["score"]:7.1f}')


if __name__ == '__main__':
    main()
